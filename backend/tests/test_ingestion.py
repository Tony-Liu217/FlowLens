from __future__ import annotations

import csv
import hashlib
import io
import json
import sqlite3
import sys
import tempfile
import unittest
from datetime import datetime, time
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'.deps'))

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from jiaowopay_ingest import import_files, load_records, save_bundle
from jiaowopay_ingest.mapping import Mapping, find_header, map_header
from jiaowopay_ingest.model import ReadResult, Row, Segment
from jiaowopay_ingest.normalize import normalize, parse_date, parse_money


class IngestionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def csv(self, name='statement.csv', rows=None, delimiter=',', encoding='utf-8-sig'):
        p = self.root/name
        with p.open('w',encoding=encoding,newline='') as f:
            csv.writer(f,delimiter=delimiter).writerows(rows or [
                ['交易日期','金额','收/支','币种','交易对方','流水号'],
                ['2026-09-01','123.45','支出','CNY','合成商户','00012345678901234567890'],
            ])
        return p

    def codes(self, d):
        return {p['code'] for p in d['issues']}

    def test_basic_read_and_no_dedupe(self):
        p = self.csv(rows=[['交易日期','金额','收/支','币种'],['2026-09-01','100','支出','CNY'],['2026-09-01','100','支出','CNY']])
        d = import_files([p])
        self.assertEqual(d['summary']['ready_count'],2)
        self.assertNotEqual(d['records'][0]['record_id'],d['records'][1]['record_id'])
        self.assertFalse(d['summary']['transaction_deduplication'])
        self.assertNotIn('category_l1',d['records'][0])

    def test_input_order_and_replay_stable_ids(self):
        a,b = self.csv(),self.csv('other.csv',rows=[['交易日期','金额','收/支','币种'],['2026-09-02','20','收入','CNY']])
        one,two=import_files([a,b]),import_files([b,a])
        self.assertEqual({r['record_id'] for r in one['records']},{r['record_id'] for r in two['records']})
        self.assertNotEqual(one['batch_id'],two['batch_id'])

    def test_file_duplicate_is_explicit_and_opt_in(self):
        p=self.csv()
        q=self.root/'renamed.csv'; q.write_bytes(p.read_bytes())
        d=import_files([p,q])
        self.assertEqual(len(d['records']),1)
        self.assertIn('DUPLICATE_FILE',self.codes(d))
        d=import_files([p,q],allow_duplicate_files=True)
        self.assertEqual(len(d['records']),2)
        self.assertEqual(len({r['raw_id'] for r in d['records']}),2)
        save_bundle(d,self.root/'dup')

    def test_split_income_expense(self):
        p=self.csv(rows=[['交易日期','收入金额','支出金额','币种'],['2026-09-01','','100','CNY'],['2026-09-02','25','','CNY']])
        d=import_files([p]); self.assertEqual(d['status'],'completed')
        self.assertEqual([(r['amount'],r['direction']) for r in d['records']],[('100','out'),('25','in')])

    def test_both_amounts_not_netted(self):
        d=import_files([self.csv(rows=[['交易日期','收入金额','支出金额','币种'],['2026-09-01','80','100','CNY']])])
        self.assertIsNone(d['records'][0]['amount'])
        self.assertIn('BOTH_DIRECTIONS',self.codes(d))

    def test_missing_fields_not_fabricated(self):
        d=import_files([self.csv(rows=[['交易日期','金额','币种'],['bad-date','bad-amount','CNY']])])
        r=d['records'][0]
        self.assertIsNone(r['transaction_at']); self.assertIsNone(r['amount']); self.assertIsNone(r['direction'])
        self.assertEqual(r['parse_status'],'needs_review')

    def test_unknown_positive_direction_is_not_income(self):
        d=import_files([self.csv(rows=[['交易日期','金额','币种'],['2026-09-01','100','CNY']])])
        self.assertIsNone(d['records'][0]['direction'])
        self.assertIn('DIRECTION_MISSING',self.codes(d))

    def test_borrow_lend_needs_account_perspective(self):
        d=import_files([self.csv(rows=[['交易日期','金额','借贷标志','币种'],['2026-09-01','100','借','CNY']])])
        self.assertEqual(d['records'][0]['parse_status'],'needs_review')

    def test_sequence_never_guessed_as_money(self):
        d=import_files([self.csv(rows=[['交易日期','序号','值'],['2026-09-01','1','100']])])
        self.assertEqual(d['records'],[])
        self.assertEqual(len(d['raw_rows']),2)
        self.assertIn('HEADER_UNRECOGNIZED',self.codes(d))

    def test_exact_counterparty_mapping(self):
        m,_=map_header(['交易日期','金额','对方账号','商户名称'])
        self.assertEqual(m['counterparty'],3)

    def test_header_ambiguity_stays_visible(self):
        d=import_files([self.csv(rows=[['交易日期','金额','交易金额','币种','收支'],['2026-09-01','10','20','CNY','支出']])])
        self.assertIn('COLUMN_AMBIGUOUS',self.codes(d))
        self.assertEqual(d['records'][0]['parse_status'],'needs_review')

    def test_preamble_row_number_and_multiline_csv(self):
        p=self.csv(rows=[['合成账单说明'],[],['交易日期','金额','收支','币种','商品说明'],['2026-09-01','10','支出','CNY','两行\n文字,带逗号']])
        d=import_files([p]);r=d['records'][0]
        self.assertEqual(r['row'],4);self.assertEqual(r['end_row'],5)
        self.assertEqual(r['description'],'两行\n文字,带逗号')
        self.assertEqual(len(d['raw_rows']),4)

    def test_repeated_header_footer_not_transactions(self):
        h=['交易日期','金额','收支','币种']
        d=import_files([self.csv(rows=[h,['2026-09-01','10','支出','CNY'],h,['2026-09-02','20','支出','CNY'],['合计','30']])])
        self.assertEqual(len(d['records']),2)
        self.assertEqual(d['raw_rows'][-1]['disposition'],'footer')

    def test_encodings_and_delimiters(self):
        for encoding,delimiter in [('utf-16','\t'),('gb18030',';'),('utf-8',',')]:
            with self.subTest(encoding=encoding):
                d=import_files([self.csv(encoding=encoding,delimiter=delimiter)])
                self.assertEqual(d['summary']['ready_count'],1)

    def test_exact_decimals_and_invalid_numeric_forms(self):
        self.assertEqual(parse_money('1,234.56'),Decimal('1234.56'))
        self.assertEqual(parse_money('（100.50）'),Decimal('-100.50'))
        for value in ('NaN','Infinity','1.234,56','1,23','1e1000',True):
            self.assertIsNone(parse_money(value))

    def test_compact_date_and_day_precision(self):
        self.assertEqual(parse_date(20260901),('2026-09-01','day'))
        self.assertEqual(parse_date('2026-09-01garbage'),(None,None))
        self.assertEqual(parse_date(46266),(None,None))

    def test_xlsx_all_sheets_dates_times_and_epoch(self):
        p=self.root/'multi.xlsx';w=Workbook();w.epoch=CALENDAR_MAC_1904
        w.active.title='封面';w.active.append(['合成文件，仅用于测试'])
        s=w.create_sheet('交易');s.append(['交易日期','交易时间','金额','收支','币种','流水号'])
        s.append([datetime(2026,9,1),time(12,30),12.30,'支出','CNY','00012345678901234567890'])
        s['A2'].number_format='yyyy-mm-dd';s['B2'].number_format='hh:mm:ss'
        w.save(p);w.close()
        d=import_files([p]);r=d['records'][0]
        self.assertEqual(r['sheet'],'交易');self.assertEqual(r['transaction_at'],'2026-09-01 12:30:00')
        self.assertEqual(r['source_record_id'],'00012345678901234567890')

    def test_excel_serial_time_without_style(self):
        h=['交易日期','时间','金额','收支','币种']
        m=find_header([Row(h,1)])
        r,_=normalize(Row([20260901,0.5,100,'支出','CNY'],2),m)
        self.assertEqual(r['transaction_at'],'2026-09-01 12:00:00')

    def test_numeric_identifiers_and_formula_cells_flagged(self):
        p=self.root/'numeric.xlsx';w=Workbook();s=w.active
        s.append(['交易日期','金额','收支','币种','流水号'])
        s.append(['2026-09-01','=10+20','支出','CNY',1234567890123456789])
        w.save(p);w.close()
        d=import_files([p])
        self.assertIn('FORMULA_CELL',self.codes(d));self.assertIn('IDENTIFIER_NUMERIC',self.codes(d))
        self.assertEqual(d['records'][0]['parse_status'],'needs_review')

    def test_closed_transaction_is_preserved_not_deleted(self):
        d=import_files([self.csv(rows=[['交易日期','金额','收支','币种','交易状态'],['2026-09-01','10','支出','CNY','交易关闭']])])
        self.assertEqual(d['records'][0]['transaction_status'],'void')
        self.assertEqual(len(d['records']),1)

    def test_user_mapping_reprocess_without_source_mutation(self):
        p=self.csv(rows=[['发生日','收入值','支出值'],['2026-09-01','','20']]);before=p.read_bytes()
        sha=hashlib.sha256(before).hexdigest()
        cfg={sha:{'CSV':{'header_row':1,'columns':{'transaction_at':0,'credit_amount':1,'debit_amount':2},'currency':'CNY'}}}
        d=import_files([p],mappings=cfg)
        self.assertEqual(d['summary']['ready_count'],1);self.assertEqual(p.read_bytes(),before)
        self.assertEqual(d['records'][0]['mapping_method'],'user_confirmed')

    def test_invalid_mapping_does_not_fallback_silently(self):
        p=self.csv();sha=hashlib.sha256(p.read_bytes()).hexdigest()
        d=import_files([p],mappings={sha:{'CSV':{'header_row':1,'columns':{'amount':999}}}})
        self.assertIn('MAPPING_INVALID',self.codes(d));self.assertEqual(d['records'],[])

    def test_pdf_continuation_and_scan_explicit(self):
        p=self.root/'sample.pdf';p.write_bytes(b'%PDF-synthetic')
        header=['交易日期','金额','收支','币种']
        result=ReadResult('pdf',[
            Segment('p1',[Row(header,1),Row(['2026-09-01','10','支出','CNY'],2)],'pdf_table',1,1,(0,20,40,60)),
            Segment('p2',[Row(['2026-09-02','20','支出','CNY'],1)],'pdf_table',2,1,(0,20,40,60)),
        ])
        with patch('jiaowopay_ingest.pipeline.read_document',return_value=result):d=import_files([p])
        self.assertEqual(len(d['records']),2);self.assertIn('HEADER_INHERITED',self.codes(d))
        result.segments[1].layout=(0,30,40,60)
        with patch('jiaowopay_ingest.pipeline.read_document',return_value=result):d=import_files([p])
        self.assertEqual(len(d['records']),1);self.assertIn('HEADER_UNRECOGNIZED',self.codes(d))

    def test_bad_file_does_not_abort_good_file(self):
        d=import_files([self.root/'missing.xls',self.csv()])
        self.assertEqual(len(d['records']),1);self.assertEqual(d['status'],'needs_review')

    def test_storage_roundtrip_and_partial_gate(self):
        d=import_files([self.csv()]);target=save_bundle(d,self.root/'out')
        self.assertEqual(load_records(target),d['records'])
        with self.assertRaises(FileExistsError):save_bundle(d,target)
        partial=import_files([self.csv('partial.csv',rows=[['交易日期','金额'],['bad','100']])])
        target=save_bundle(partial,self.root/'partial',excel=False)
        with self.assertRaises(ValueError):load_records(target)
        self.assertEqual(load_records(target,allow_partial=True),[])
        self.assertEqual(len(load_records(target,allow_partial=True,ready_only=False)),1)

    def test_export_no_formulas_in_source_text(self):
        p=self.csv(rows=[['交易日期','金额','收支','币种','商品说明'],['2026-09-01','10','支出','CNY','=HYPERLINK("https://invalid.example", "x")']])
        d=import_files([p]);target=save_bundle(d,self.root/'export')
        w=load_workbook(target/'标准化流水.xlsx',data_only=False)
        self.assertEqual(w.sheetnames,['标准化流水','原始字段','异常与待确认','导入说明'])
        self.assertEqual(w['标准化流水']['I2'].data_type,'s')
        self.assertEqual(w['标准化流水']['D2'].value,10)
        w.close()

    def test_no_network_even_if_key_exists(self):
        with patch.dict('os.environ',{'DEEPSEEK_API_KEY':'test-only-not-a-secret'}),patch('socket.socket',side_effect=AssertionError('network forbidden')):
            d=import_files([self.csv()])
        self.assertEqual(d['status'],'completed');self.assertFalse(d['network_used'])

    def test_two_line_header(self):
        p=self.csv(rows=[['交易日期','收入','支出','币种'],['','金额','金额',''],['2026-09-01','','20','CNY']])
        d=import_files([p]);self.assertEqual(len(d['records']),1)
        self.assertEqual(d['records'][0]['amount'],'20')

    def test_same_filename_different_contents_kept(self):
        a=self.csv()
        folder=self.root/'another';folder.mkdir()
        b=folder/a.name
        b.write_text('交易日期,金额,收支,币种\n2026-09-02,99,支出,CNY\n',encoding='utf-8')
        d=import_files([a,b])
        self.assertEqual(len(d['records']),2)
        self.assertEqual(len({r['file_id'] for r in d['records']}),2)

    def test_actual_content_overrides_extension(self):
        p=self.root/'misnamed.csv'
        w=Workbook();s=w.active
        s.append(['交易日期','金额','收支','币种']);s.append(['2026-09-01',10,'支出','CNY'])
        w.save(p);w.close()
        d=import_files([p])
        self.assertEqual(d['summary']['ready_count'],1)
        self.assertIn('EXTENSION_MISMATCH',self.codes(d))

    def test_missing_currency_defaults_cny_without_review(self):
        d=import_files([self.csv(rows=[['交易日期','金额','收支'],['2026-09-01','10','支出']])])
        self.assertEqual(d['records'][0]['currency'], 'CNY')
        self.assertEqual(d['records'][0]['parse_status'], 'ready')
        self.assertIn('CURRENCY_DEFAULTED',self.codes(d))
        self.assertEqual(next(i for i in d['issues'] if i['code']=='CURRENCY_DEFAULTED')['severity'], 'warning')

    def test_unknown_currency_defaults_but_foreign_currency_is_preserved(self):
        for raw, expected in [('?', 'CNY'), ('USD', 'USD'), ('美元', 'USD'), ('EUR', 'EUR')]:
            with self.subTest(raw=raw):
                d=import_files([self.csv(rows=[['交易日期','金额','收支','币种'],['2026-09-01','10','支出',raw]])])
                self.assertEqual(d['records'][0]['currency'], expected)
                self.assertEqual(d['records'][0]['parse_status'], 'ready')
                target=save_bundle(d,self.root/('currency-'+raw.replace('?', 'unknown')),excel=False)
                self.assertEqual(load_records(target)[0]['currency'],expected)

    def test_default_currency_does_not_hide_invalid_amount(self):
        d=import_files([self.csv(rows=[['交易日期','金额','收支'],['2026-09-01','xx','支出']])])
        self.assertEqual(d['records'][0]['currency'], 'CNY')
        self.assertEqual(d['records'][0]['parse_status'], 'needs_review')
        self.assertIn('AMOUNT_INVALID',self.codes(d))

    def test_signed_amount_direction_conflict(self):
        d=import_files([self.csv(rows=[['交易日期','金额','收支','币种'],['2026-09-01','-10','收入','CNY']])])
        self.assertIn('DIRECTION_CONFLICT',self.codes(d))
        self.assertEqual(d['records'][0]['parse_status'],'needs_review')

    def test_pdf_text_outside_table_does_not_disappear(self):
        p=self.root/'text.pdf';p.write_bytes(b'%PDF-synthetic')
        result=ReadResult('pdf',[
            Segment('p1',[Row(['交易日期','金额','收支','币种'],1),Row(['2026-09-01','10','支出','CNY'],2)],'pdf_table',1,1,(0,20,40,60)),
            Segment('p1-text',[Row(['2026-09-02 合成商户 -200.00 余额800.00'],1)],'pdf_text',1),
        ])
        with patch('jiaowopay_ingest.pipeline.read_document',return_value=result):d=import_files([p])
        self.assertIn('UNMAPPED_TEXT',self.codes(d))
        self.assertEqual(len(d['records']),1)
        self.assertEqual(d['raw_rows'][-1]['disposition'],'document_text')

    def test_empty_and_corrupt_files_are_not_success(self):
        a=self.root/'empty.csv';a.write_bytes(b'')
        b=self.root/'bad.pdf';b.write_bytes(b'not a pdf')
        d=import_files([a,b])
        self.assertEqual(d['status'],'needs_review')
        self.assertIn('NO_TRANSACTIONS',self.codes(d))
        self.assertIn('FORMAT_MISMATCH',self.codes(d))

    def test_decimal_roundtrip_and_no_cross_currency_sum(self):
        p=self.csv(rows=[['交易日期','金额','收支','币种'],['2026-09-01','0.10','支出','CNY'],['2026-09-02','0.20','支出','USD']])
        d=import_files([p]);target=save_bundle(d,self.root/'decimal',excel=False)
        rows=load_records(target)
        self.assertEqual([r['amount'] for r in rows],['0.10','0.20'])
        self.assertNotIn('total_amount',d['summary'])

    def test_incomplete_and_wrong_schema_rejected(self):
        p=self.root/'incomplete';p.mkdir()
        with self.assertRaises(ValueError):load_records(p)
        d=import_files([self.csv()]);target=save_bundle(d,self.root/'schema',excel=False)
        con=sqlite3.connect(target/'ledger.sqlite3')
        con.execute('UPDATE manifest SET value=? WHERE key=?',(json.dumps('future'),'schema_version'))
        con.commit();con.close()
        with self.assertRaises(ValueError):load_records(target)

    def test_nan_dates_and_ids_do_not_crash(self):
        self.assertEqual(parse_date(float('nan')),(None,None))
        m=find_header([Row(['交易日期','金额','收支','币种','流水号'],1)])
        r,problems=normalize(Row(['2026-09-01','10','支出','CNY',float('nan')],2),m)
        self.assertEqual(r['parse_status'],'needs_review')
        self.assertIn('IDENTIFIER_NUMERIC',{p['code'] for p in problems})

    def test_original_files_not_changed(self):
        p=self.csv();before=p.read_bytes()
        import_files([p]);self.assertEqual(p.read_bytes(),before)

    def test_zero_is_distinct_from_missing(self):
        d=import_files([self.csv(rows=[['交易日期','金额','收支','币种'],['2026-09-01','0','不计收支','CNY']])])
        self.assertEqual(d['records'][0]['amount'],'0')
        self.assertEqual(d['records'][0]['parse_status'],'ready')

    def test_user_direction_code_is_explicit(self):
        p=self.csv(rows=[['交易日期','金额','借贷标志'],['2026-09-01','10','借']])
        sha=hashlib.sha256(p.read_bytes()).hexdigest()
        config={sha:{'CSV':{'header_row':1,'columns':{'transaction_at':0,'amount':1,'direction':2},'currency':'CNY','direction_codes':{'借':'out','贷':'in'}}}}
        d=import_files([p],mappings=config)
        self.assertEqual(d['records'][0]['direction'],'out')
        self.assertEqual(d['records'][0]['parse_status'],'ready')


if __name__=='__main__':
    unittest.main()
