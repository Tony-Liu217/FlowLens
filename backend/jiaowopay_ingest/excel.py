"""Application export adapter. Workbook is a review view, not a round-trip database."""
from __future__ import annotations

import re
import math
from datetime import datetime
from decimal import Decimal
from weakref import WeakKeyDictionary

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

from .model import dumps

STATUS_LABELS = {"ready": "已识别", "needs_review": "待确认", "in": "收入", "out": "支出", "none": "不计收支", "posted": "已完成", "void": "关闭/失败", "pending": "处理中", "unknown": "未知"}
_row_counts = WeakKeyDictionary()
COLUMNS = [
    ("parse_status", "解析状态", 14), ("transaction_at", "交易日期/时间", 23),
    ("booking_at", "记账日期/时间", 23), ("amount", "金额（原币）", 18),
    ("currency", "币种", 10), ("direction", "收支方向", 14), ("source", "账单来源", 16),
    ("counterparty", "交易对方", 28), ("description", "商品/摘要", 40),
    ("raw_status", "原始状态", 20), ("transaction_status", "标准状态", 14),
    ("source_record_id", "原始流水号", 34), ("merchant_order_no", "商户订单号", 34),
    ("trade_type", "原始交易类型", 20), ("original_category", "原始分类", 18),
    ("payment_method", "支付方式/渠道", 24), ("memo", "备注/附言", 38),
    ("balance", "余额（原币）", 18), ("counterparty_account", "对方账号", 30),
    ("counterparty_combined", "对方账号与户名原文", 36), ("counterparty_bank", "对方开户行", 30),
    ("transaction_at_precision", "交易时间精度", 18), ("booking_at_precision", "记账时间精度", 18),
    ("filename", "原始文件", 42), ("sheet", "工作表", 18), ("page", "页码", 10),
    ("table", "表格编号", 12), ("row", "原始行号", 12), ("profile", "适配规则", 28),
    ("record_id", "记录标识", 68), ("raw_id", "原始行标识", 68),
    ('review_id','人工核验任务标识',68),
]


def _append(ws, values):
    values = [v[:30000]+" [显示节选；完整内容见原始字段或内部数据]" if isinstance(v,str) and len(v)>32700 else v for v in values]
    ws.append(values)
    row_number = _row_counts.get(ws, 0) + 1
    _row_counts[ws] = row_number
    for column in range(1, len(values)+1):
        cell = ws.cell(row_number, column)
        if isinstance(cell.value, str):
            cell.value = ILLEGAL_CHARACTERS_RE.sub(lambda m: "\\u%04x" % ord(m[0]), cell.value)
            # Never turn source strings into Excel formulas or automatic links.
            cell.data_type = "s"
    return row_number


def _style(ws, widths):
    ws.freeze_panes = "B2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 28
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="20354D")
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i,w in enumerate(widths,1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        lines = 1
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            width = widths[cell.column-1]
            needed = sum(max(1, math.ceil(sum(2 if ord(c)>255 else 1 for c in part)/max(width-2,1))) for part in cell.value.splitlines())
            if needed > 1:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                lines = max(lines, needed)
        if lines > 1:
            ws.row_dimensions[row[0].row].height = min(150,16*lines)


def export_xlsx(dataset, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "标准化流水"
    _append(ws, [c[1] for c in COLUMNS])
    for record in dataset["records"]:
        values = []
        for key,_,_ in COLUMNS:
            value = record.get(key)
            if key in {"parse_status", "direction", "transaction_status"}:
                value = STATUS_LABELS.get(value,value)
            elif key in {"transaction_at", "booking_at"} and value:
                value = datetime.fromisoformat(value)
            elif key in {"amount", "balance"} and value is not None:
                number = Decimal(value)
                if len(number.as_tuple().digits) <= 15:
                    value = number
            values.append(value)
        row_no = _append(ws, values)
        for i,key in ((2,"transaction_at"),(3,"booking_at")):
            ws.cell(row_no,i).number_format = {"day":"yyyy-mm-dd", "minute":"yyyy-mm-dd hh:mm"}.get(record.get(key+"_precision"), "yyyy-mm-dd hh:mm:ss")
        for i in (4,18):
            value = record.get("amount" if i == 4 else "balance")
            decimals = max(2, -Decimal(value).as_tuple().exponent) if value is not None else 2
            pattern = '#,##0.' + '0' * min(decimals, 12)
            ws.cell(row_no,i).number_format = pattern + ';[Red](' + pattern + ')'
    _style(ws,[c[2] for c in COLUMNS])
    ws.sheet_properties.tabColor = "20354D"
    if ws.max_row > 1:
        ws.conditional_formatting.add(f"A2:A{ws.max_row}", CellIsRule(operator="equal", formula=['"待确认"'], fill=PatternFill("solid",fgColor="FFF0CE")))

    raw_sheet = wb.create_sheet("原始字段")
    _append(raw_sheet,["原始行标识","原始文件","片段/工作表","页码","表格编号","原始行号","行去向","原始列号","原始字段名","值类型","原始值","内容分段","源公式"])
    records_by_raw = {r["raw_id"]:r for r in dataset["records"]}
    for raw in dataset["raw_rows"]:
        record = records_by_raw.get(raw["raw_id"],{})
        headers = record.get("raw_headers",[])
        for i,value in enumerate(raw["values"]):
            if value is None or value == "":
                continue
            kind = value.get("type") if isinstance(value,dict) else type(value).__name__
            string = value.get("value") if isinstance(value,dict) else str(value)
            chunks = [string[j:j+30000] for j in range(0,len(string),30000)] or [""]
            for n,chunk in enumerate(chunks,1):
                _append(raw_sheet,[raw["raw_id"],raw["filename"],raw["segment"],raw["page"],raw["table"],raw["row"],raw["disposition"],i+1,text_header(headers,i),kind,chunk,f"{n}/{len(chunks)}",raw["formulas"].get(i)])
    _style(raw_sheet,[68,42,28,10,12,12,18,12,24,14,60,12,40])

    problems = wb.create_sheet("异常与待确认")
    _append(problems,["级别","问题代码","说明","字段","原始文件","片段/工作表","页码","原始行标识","记录标识"])
    file_names = {f.get("file_id"):f["filename"] for f in dataset["files"]}
    for p in dataset["issues"]:
        _append(problems,["需处理" if p["severity"] == "error" else "提示",p["code"],p["message"],p.get("field"),file_names.get(p.get("file_id"),p.get("filename")),p.get("segment"),p.get("page"),p.get("raw_id"),p.get("record_id")])
    _style(problems,[12,30,88,24,42,24,10,68,68])

    info = wb.create_sheet("导入说明")
    _append(info,["项目","内容","补充说明"])
    for label,value in [("用途","未去重的标准化流水，非实际花费报告"),("状态",dataset["status"]),("批次",dataset["batch_id"]),("数据版本",dataset["schema_version"]),("解析版本",dataset["parser_version"]),("外部网络调用","无"),("原始文件","未复制原文件，请自行保留；内部数据保留已提取的原始值和位置"),("注意","已识别仅代表字段通过解析，不代表交易已完成或应计入消费"),("后续读取","使用 ledger.sqlite3 与 load_records；修改 Excel 不会改动内部数据"),("敏感数据","此数据包包含账单原文、账号等信息，只用于本地核对，勿公开")]:
        _append(info,[label,value,None])
    for key,value in dataset["summary"].items():
        _append(info,[key,value,None])
    for f in dataset["files"]:
        _append(info,[f["filename"],f["status"],f"候选记录 {f.get('record_count',0)}，已识别 {f.get('ready_count',0)}"])
    _style(info,[46,100,65])
    wb.save(path)
    wb.close()


def text_header(headers,index):
    value = headers[index] if index < len(headers) else None
    return str(value) if value is not None else None
