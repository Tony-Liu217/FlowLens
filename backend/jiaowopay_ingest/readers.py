"""Physical readers only: preserve row positions, types, page and table boundaries."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

from .model import ReadFailure, ReadResult, Row, Segment, issue

MAX_BYTES = 50 * 1024 * 1024
MAX_EXPANDED_BYTES = 250 * 1024 * 1024
MAX_ROWS = 200_000
MAX_COLUMNS = 256
MAX_PAGES = 300


def detect_format(path: Path) -> str:
    if path.stat().st_size > MAX_BYTES:
        raise ReadFailure("FILE_TOO_LARGE", "单文件超过 50 MiB，未读取。")
    with path.open("rb") as stream:
        magic = stream.read(1024)
    if magic.startswith(b"%PDF-"):
        return "pdf"
    if magic.startswith(b'\x89PNG\r\n\x1a\n'):
        return 'png'
    if magic.startswith(b'\xff\xd8\xff'):
        return 'jpeg'
    if magic.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls"
    if magic.startswith(b"PK"):
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                    raise ReadFailure("ARCHIVE_TOO_LARGE", "工作簿解压体积或条目数超过限制。")
                if "xl/workbook.xml" in archive.namelist():
                    return "xlsx"
        except zipfile.BadZipFile as exc:
            raise ReadFailure("CORRUPT_WORKBOOK", "工作簿容器损坏。") from exc
        raise ReadFailure("UNSUPPORTED_ARCHIVE", "这不是支持的 Excel 工作簿；请解压账单后分别导入。")
    if path.suffix.lower() in {".csv", ".tsv", ".txt"}:
        return "csv"
    raise ReadFailure("FORMAT_MISMATCH", "文件内容不符合支持的 CSV、XLSX、XLS、PDF、PNG 或 JPEG 格式。")


def _csv(path: Path) -> ReadResult:
    data = path.read_bytes()
    encodings = ["utf-8-sig", "gb18030"]
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings = ["utf-16"]
    text, encoding = None, None
    for enc in encodings:
        try:
            text, encoding = data.decode(enc, errors="strict"), enc
            break
        except UnicodeDecodeError:
            continue
    if text is None or "\x00" in text:
        raise ReadFailure("ENCODING_UNKNOWN", "编码无法可靠识别，请导出 UTF-8 CSV 或带 BOM 的 UTF-16 文件。")
    # A preamble must not determine the dialect. Prefer repeatedly structured rows.
    best, best_score = ",", -1
    sample = text[:200000]
    for delimiter in (",", "\t", ";", "，"):
        try:
            rows = list(csv.reader(io.StringIO(sample), delimiter=delimiter))
        except csv.Error:
            continue
        score = sum(1 for r in rows if len(r) >= 3)
        if score > best_score:
            best, best_score = delimiter, score
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=best, strict=True)
    rows, previous = [], 0
    try:
        for values in reader:
            if len(values) > MAX_COLUMNS or len(rows) >= MAX_ROWS:
                raise ReadFailure("TABLE_TOO_LARGE", "表格行列数超过导入限制。")
            rows.append(Row(values, previous + 1, reader.line_num, ["text"] * len(values)))
            previous = reader.line_num
    except csv.Error as exc:
        raise ReadFailure("CSV_SYNTAX", f"CSV 在物理行 {reader.line_num} 附近引号或字段结构不完整。") from exc
    return ReadResult("csv", [Segment("CSV", rows)], metadata={"encoding": encoding, "delimiter": best})


def _xlsx(path: Path) -> ReadResult:
    import openpyxl

    result = ReadResult("xlsx")
    with path.open("rb") as stream, path.open("rb") as formula_stream:
        workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True, keep_links=False)
        formulas = openpyxl.load_workbook(formula_stream, read_only=True, data_only=False, keep_links=False)
        try:
            result.metadata["excel_epoch"] = workbook.epoch.isoformat()
            total = 0
            for sheet, formula_sheet in zip(workbook, formulas):
                # Read-only dimensions can be stale. Iterate actual XML rows.
                sheet.reset_dimensions()
                formula_sheet.reset_dimensions()
                rows = []
                result.segments.append(Segment(sheet.title, rows))
                for row_no, (cells, formula_cells) in enumerate(zip(sheet.iter_rows(), formula_sheet.iter_rows()), 1):
                    total += 1
                    if total > MAX_ROWS or len(cells) > MAX_COLUMNS:
                        raise ReadFailure("TABLE_TOO_LARGE", "工作簿行列数超过导入限制。")
                    values = []
                    for c in cells:
                        value = c.value
                        if isinstance(value, datetime) and not re.search(r"[hs]", c.number_format, re.I):
                            value = value.date()
                        values.append(value)
                    rows.append(Row(
                        values, row_no,
                        cell_types=[c.data_type for c in cells],
                        formats=[c.number_format for c in cells],
                        formulas={i: c.value for i, c in enumerate(formula_cells) if c.data_type == "f"},
                    ))
        finally:
            workbook.close()
            formulas.close()
    return result


def _xls(path: Path) -> ReadResult:
    import xlrd

    book = xlrd.open_workbook(path, formatting_info=True, on_demand=True)
    result = ReadResult("xls", metadata={"excel_datemode": book.datemode})
    total = 0
    try:
        for sheet in book.sheets():
            total += sheet.nrows
            if total > MAX_ROWS or sheet.ncols > MAX_COLUMNS:
                raise ReadFailure("TABLE_TOO_LARGE", "工作簿行列数超过导入限制。")
            rows = []
            for r in range(sheet.nrows):
                values, types, formats = [], [], []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                        if 0 <= cell.value < 1:
                            value = value.time()
                    fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                    if isinstance(value, datetime) and not re.search(r"[hs]", fmt, re.I):
                        value = value.date()
                    values.append(value)
                    types.append({xlrd.XL_CELL_NUMBER: "n", xlrd.XL_CELL_DATE: "d", xlrd.XL_CELL_ERROR: "e"}.get(cell.ctype, "text"))
                    formats.append(fmt)
                rows.append(Row(values, r+1, cell_types=types, formats=formats))
            result.segments.append(Segment(sheet.name, rows))
    finally:
        book.release_resources()
    return result


def _pdf(path: Path) -> ReadResult:
    import pdfplumber

    result = ReadResult("pdf")
    with pdfplumber.open(path) as pdf:
        if len(pdf.pages) > MAX_PAGES:
            raise ReadFailure("TOO_MANY_PAGES", "PDF 超过 300 页，请按账期分批导入。")
        result.metadata["page_count"] = len(pdf.pages)
        total = 0
        for page in pdf.pages:
            try:
                if not page.chars:
                    result.issues.append(issue("OCR_REQUIRED", "此页没有可提取文字，需要本地 OCR 或重新导出文本账单。", page=page.page_number))
                    result.segments.append(Segment(f"page-{page.page_number}", [], "scan", page.page_number))
                    continue
                tables = page.find_tables()
                # Includes scanned pages with text-only titles/watermarks and
                # text PDFs without extractable tables. Route the whole page
                # once; do not combine it with overlapping native text records.
                if not tables:
                    result.issues.append(issue('OCR_REQUIRED','此页无法提取结构化表格，转本地 OCR；不会与文字记录重复导入。',page=page.page_number))
                    result.segments.append(Segment(f'page-{page.page_number}',[],'scan',page.page_number))
                    continue
                bounds = []
                for i, table in enumerate(tables, 1):
                    values = table.extract()
                    if not any(any(v not in (None, "") for v in row) for row in values):
                        continue
                    total += len(values)
                    if total > MAX_ROWS or max(map(len, values), default=0) > MAX_COLUMNS:
                        raise ReadFailure("TABLE_TOO_LARGE", "PDF 表格行列数超过限制。")
                    # Column geometry protects continuation mapping against reordered layouts.
                    xs = sorted({round(cell[0], 0) for cell in table.cells if cell is not None})
                    result.segments.append(Segment(f"page-{page.page_number}-table-{i}", [Row(list(v), j+1) for j,v in enumerate(values)], "pdf_table", page.page_number, i, tuple(xs)))
                    bounds.append(table.bbox)
                outside = page.filter(lambda o: not any(b[0] <= o.get("x0", -1) <= b[2] and b[1] <= o.get("top", -1) <= b[3] for b in bounds))
                lines = (outside.extract_text() or "").splitlines()
                if lines:
                    result.segments.append(Segment(f"page-{page.page_number}-text", [Row([v], i+1) for i,v in enumerate(lines)], "pdf_text", page.page_number))
                if not tables:
                    result.issues.append(issue("PDF_LAYOUT_UNSUPPORTED", "此页未提取出表格；原文已保留，未猜测交易金额。", page=page.page_number))
            except ReadFailure:
                raise
            except Exception as exc:
                result.issues.append(issue("PDF_PAGE_FAILED", f"此页提取失败（{type(exc).__name__}）。", page=page.page_number))
    return result


def read_document(path: Path) -> ReadResult:
    fmt = detect_format(path)
    try:
        if fmt in {'png','jpeg'}:
            result=ReadResult(fmt,[Segment('image-1',[],'scan',1)],
                              [issue('OCR_REQUIRED','图片需要本地 OCR。',page=1)])
        else:
            result = {"csv": _csv, "xlsx": _xlsx, "xls": _xls, "pdf": _pdf}[fmt](path)
    except ImportError as exc:
        raise ReadFailure("DEPENDENCY_MISSING", "缺少读取依赖，请安装 backend/requirements.txt。") from exc
    except ReadFailure:
        raise
    except Exception as exc:
        # Do not echo data-bearing parser exceptions into terminal logs.
        raise ReadFailure("DOCUMENT_UNREADABLE", f"文件损坏、加密或结构不支持（{type(exc).__name__}）。") from exc
    if path.suffix.lower().lstrip(".") not in ({'jpg','jpeg'} if fmt=='jpeg' else {fmt} if fmt != "csv" else {"csv", "tsv", "txt"}):
        result.issues.append(issue("EXTENSION_MISMATCH", f"扩展名与内容不一致，已按 {fmt.upper()} 读取。", "warning"))
    return result
